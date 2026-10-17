import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app


class CareerForgeApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.data_path = Path(self.temporary_directory.name)
        self.data_patch = patch("backend.database.DATA_DIR", self.data_path)
        self.database_patch = patch(
            "backend.database.DATABASE_PATH", self.data_path / "tests.db"
        )
        self.data_patch.start()
        self.database_patch.start()
        self.client_context = TestClient(app)
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        self.database_patch.stop()
        self.data_patch.stop()
        self.temporary_directory.cleanup()

    @staticmethod
    def pdf_bytes() -> bytes:
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text(
            (72, 72),
            "Jordan Example\njordan@example.test\nSUMMARY\nAnalyst with Python skills.\n"
            "EDUCATION\nBSc Statistics\nSKILLS\nPython, SQL\nPROJECTS\nBuilt a dashboard for 20 users",
        )
        content = document.tobytes()
        document.close()
        return content

    @staticmethod
    def docx_bytes() -> bytes:
        document = Document()
        document.add_paragraph("Jordan Example")
        document.add_paragraph("jordan@example.test")
        document.add_heading("Skills", level=2)
        document.add_paragraph("Python, SQL")
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def test_health_does_not_disclose_credentials(self) -> None:
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        self.assertIn("ai_configured", response.json())
        self.assertNotIn("api_key", response.json())

    def test_uploads_pdf_and_returns_extracted_profile(self) -> None:
        response = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", self.pdf_bytes(), "application/pdf")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["profile"]["email"], "jordan@example.test")
        self.assertIn("Python", response.json()["profile"]["skills"])

    def test_uploads_docx(self) -> None:
        response = self.client.post(
            "/api/resumes/upload",
            files={
                "file": (
                    "resume.docx",
                    self.docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("SQL", response.json()["profile"]["skills"])

    def test_rejects_unsupported_empty_and_corrupt_files(self) -> None:
        unsupported = self.client.post(
            "/api/resumes/upload", files={"file": ("resume.txt", b"text", "text/plain")}
        )
        empty = self.client.post(
            "/api/resumes/upload", files={"file": ("resume.pdf", b"", "application/pdf")}
        )
        corrupt = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", b"not a pdf", "application/pdf")},
        )
        self.assertEqual(unsupported.status_code, 415)
        self.assertEqual(empty.status_code, 400)
        self.assertEqual(corrupt.status_code, 422)

    def test_heuristic_review_match_and_delete_workflow(self) -> None:
        uploaded = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", self.pdf_bytes(), "application/pdf")},
        )
        resume_id = uploaded.json()["id"]
        analysis = self.client.post(
            f"/api/resumes/{resume_id}/analyze",
            json={"target_role": "Analyst", "use_ai": False},
        )
        match = self.client.post(
            "/api/jobs/match",
            json={
                "resume_id": resume_id,
                "job_title": "Data analyst",
                "job_description": "We need Python SQL skills and dashboard experience for our data team.",
                "use_ai": False,
            },
        )
        self.assertEqual(analysis.status_code, 200)
        self.assertEqual(analysis.json()["source"], "heuristic")
        self.assertEqual(match.status_code, 200)
        self.assertGreater(match.json()["score"], 0)
        deleted = self.client.delete(f"/api/resumes/{resume_id}")
        self.assertEqual(deleted.status_code, 200)
        dashboard = self.client.get("/api/dashboard").json()
        self.assertEqual(dashboard["counts"]["resumes"], 0)
        self.assertEqual(len(dashboard["activity"]), 1)
        self.assertEqual(self.client.get(f"/api/resumes/{resume_id}").status_code, 404)

    def test_ai_request_reports_missing_credentials(self) -> None:
        uploaded = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", self.pdf_bytes(), "application/pdf")},
        )
        with patch("backend.services.gemini.GEMINI_API_KEY", ""):
            response = self.client.post(
                f"/api/resumes/{uploaded.json()['id']}/analyze",
                json={"target_role": "Analyst", "use_ai": True},
            )
        self.assertEqual(response.status_code, 503)
        self.assertIn("GEMINI_API_KEY", response.json()["detail"])

    def test_gemini_analysis_response_is_structured_and_keeps_bullet_drafts(self) -> None:
        uploaded = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", self.pdf_bytes(), "application/pdf")},
        )
        ai_result = {
            "strengths": ["The resume includes a project."],
            "suggestions": ["Clarify your personal contribution."],
            "missing_sections": ["Certifications"],
            "bullet_rewrites": [
                {
                    "original": "Built a dashboard for 20 users",
                    "suggested": "Built a dashboard serving 20 users.",
                }
            ],
            "role_alignment": ["The project relates to analysis."],
        }
        with patch(
            "backend.services.career_logic.generate_json",
            new=AsyncMock(return_value=ai_result),
        ):
            response = self.client.post(
                f"/api/resumes/{uploaded.json()['id']}/analyze",
                json={"target_role": "Data analyst", "use_ai": True},
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["source"], "gemini")
        self.assertEqual(response.json()["bullet_rewrites"], ai_result["bullet_rewrites"])
        self.assertEqual(response.json()["strengths"], ai_result["strengths"])

    def test_interview_questions_feedback_and_progress(self) -> None:
        fake_ai = AsyncMock(
            side_effect=[
                {
                    "questions": [
                        {
                            "question": "Tell me about a project you are proud of.",
                            "type": "behavioral",
                            "focus": "Your contribution",
                        }
                    ]
                },
                {
                    "scores": {
                        "relevance": 4,
                        "clarity": 3,
                        "correctness": 4,
                        "communication": 4,
                        "examples": 3,
                    },
                    "strengths": ["You included a specific project."],
                    "improvements": ["Clarify your actions."],
                    "model_answer": "A generic example answer.",
                    "star_guidance": ["Situation", "Task", "Action", "Result"],
                },
            ]
        )
        with patch("backend.services.career_logic.generate_json", new=fake_ai):
            questions = self.client.post(
                "/api/interviews/questions",
                json={"role": "Analyst", "category": "behavioral"},
            )
            feedback = self.client.post(
                "/api/interviews/feedback",
                json={
                    "question": questions.json()["questions"][0]["question"],
                    "answer": "I built a dashboard for my class project.",
                    "question_type": "behavioral",
                },
            )
        dashboard = self.client.get("/api/dashboard").json()
        self.assertEqual(questions.status_code, 200)
        self.assertEqual(feedback.status_code, 200)
        self.assertEqual(feedback.json()["scores"]["relevance"], 4)
        self.assertEqual(dashboard["counts"]["interview_answers"], 1)

    def test_roadmap_progress_is_persisted(self) -> None:
        roadmap = {
            "goal": "Practice role skills",
            "weeks": [{"id": "week-1", "title": "Foundations", "completed": True}],
        }
        updated = self.client.put("/api/roadmaps/progress", json={"roadmap": roadmap})
        dashboard = self.client.get("/api/dashboard")
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(dashboard.json()["latest_roadmap"]["weeks"][0]["completed"], True)

    def test_roadmap_generation_without_resume_is_counted(self) -> None:
        generated = {
            "goal": "Learn data analysis",
            "weeks": [{"title": "Foundations", "skills": ["SQL"], "estimated_hours": 8}],
        }
        with patch(
            "backend.routes.api.generate_roadmap",
            new=AsyncMock(return_value=generated),
        ):
            response = self.client.post(
                "/api/roadmaps",
                json={"target_role": "Data analyst", "skills": [], "missing_skills": ["SQL"]},
            )
        dashboard = self.client.get("/api/dashboard").json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(dashboard["counts"]["roadmaps"], 1)
        self.assertEqual(dashboard["latest_roadmap"]["goal"], "Learn data analysis")

    def test_profile_edits_preserve_original_and_export_pdf(self) -> None:
        uploaded = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", self.pdf_bytes(), "application/pdf")},
        )
        resume_id = uploaded.json()["id"]
        edited = self.client.put(
            f"/api/resumes/{resume_id}/profile",
            json={"name": "Jordan Example", "summary": "An edited, accurate summary.", "skills": ["Python"]},
        )
        detail = self.client.get(f"/api/resumes/{resume_id}").json()
        exported = self.client.post(
            f"/api/resumes/{resume_id}/export",
            json={"name": "Jordan Example", "summary": "An edited, accurate summary.", "skills": ["Python"]},
        )
        self.assertEqual(edited.status_code, 200)
        self.assertIn("Analyst with Python skills", detail["original_text"])
        self.assertEqual(edited.json()["profile"]["summary"], "An edited, accurate summary.")
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported.headers["content-type"], "application/pdf")
        self.assertTrue(exported.content.startswith(b"%PDF"))

    def test_delete_all_removes_unlinked_workspace_records(self) -> None:
        self.client.put(
            "/api/roadmaps/progress",
            json={"roadmap": {"goal": "Private goal", "weeks": []}},
        )
        response = self.client.delete("/api/data")
        dashboard = self.client.get("/api/dashboard").json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(dashboard["counts"]["roadmaps"], 0)
        self.assertEqual(dashboard["latest_roadmap"], None)
        self.assertEqual(dashboard["activity"], [])

    def test_upload_rejects_files_over_size_limit(self) -> None:
        response = self.client.post(
            "/api/resumes/upload",
            files={"file": ("resume.pdf", b"x" * (8 * 1024 * 1024 + 1), "application/pdf")},
        )
        self.assertEqual(response.status_code, 413)


if __name__ == "__main__":
    unittest.main()
